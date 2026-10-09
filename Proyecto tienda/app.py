from flask import Flask, render_template, request, redirect, url_for, session, flash, jsonify, send_file
import pyodbc

import logging
from datetime import datetime

logging.basicConfig(
    filename="errores.log",
    level=logging.ERROR,
    format="%(asctime)s - %(levelname)s - %(message)s"
)

def manejar_error(e, mensaje_usuario="Hubo un problema, intenta de nuevo."):
    logging.error(str(e))
    flash(mensaje_usuario, "error")

app = Flask(__name__)
app.secret_key = "pahjal_secret_2024"

SERVIDOR   = ".\SQLEXPRESS"  
BASE_DATOS = "Tienda"  
USUARIO_BD = "Usuario1"        
PASSWORD_BD= "usuario1"       

def conectar():
    return pyodbc.connect(
        f"DRIVER={{ODBC Driver 17 for SQL Server}};"
        f"SERVER={SERVIDOR};"
        f"DATABASE={BASE_DATOS};"
        f"UID={USUARIO_BD};"
        f"PWD={PASSWORD_BD};"
    )

# ─────────────────────────────────────────
#  SKU AUTOMÁTICO
# ─────────────────────────────────────────
def generar_sku():
    try:
        con = conectar()
        cursor = con.cursor()
        cursor.execute("""
            SELECT TOP 1 CODIGO_BARRAS FROM VARIANTE
            WHERE CODIGO_BARRAS LIKE 'PAH-%'
            ORDER BY id_VARIANTE DESC
        """)
        ultimo = cursor.fetchone()
        con.close()
        if ultimo and ultimo[0]:
            numero = int(ultimo[0].split("-")[1]) + 1
        else:
            numero = 1
        return f"PAH-{numero:04d}"
    except:
        return "PAH-0001"

# ═══════════════════════════════════════════
#  LOGIN
# ═══════════════════════════════════════════
@app.route("/", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        correo     = request.form.get("correo", "").strip()
        contrasena = request.form.get("contrasena", "").strip()
        if not correo or not contrasena:
            flash("Por favor ingresa correo y contraseña.", "error")
            return render_template("login.html")
        try:
            con = conectar()
            cursor = con.cursor()
            cursor.execute("""
                SELECT U.id_USUARIOS, U.NOMBRE, U.APELLID_P, U.id_SUCURSAL,U.ROL
                FROM USUARIOS U
                INNER JOIN LOGIN L ON U.id_LOGIN = L.id_LOGIN
                WHERE U.CORREO = ? AND L.CONTRASEÑA = ?
            """, (correo, contrasena))
            usuario = cursor.fetchone()
            con.close()
            if usuario:
                session["id_usuario"]  = usuario[0]
                session["nombre"]      = usuario[1]
                session["apellido"]    = usuario[2]
                session["id_sucursal"] = usuario[3]
                session["rol"]         = usuario[4]
                return redirect(url_for("bienvenida"))
            else:
                flash("Correo o contraseña incorrectos.", "error")
        except Exception as e:
            manejar_error(e, "Correo o contraseña incorrectos.")
    return render_template("login.html")

# ═══════════════════════════════════════════
#  BIENVENIDA
# ═══════════════════════════════════════════
@app.route("/bienvenida")
def bienvenida():
    if "nombre" not in session:
        return redirect(url_for("login"))
    return render_template("bienvenida.html",
                           nombre=session["nombre"],
                           apellido=session["apellido"])

# ═══════════════════════════════════════════
#  APIS CASCADA UBICACIÓN
# ═══════════════════════════════════════════
@app.route("/api/estados/<int:id_pais>")
def api_estados(id_pais):
    try:
        con = conectar()
        cursor = con.cursor()
        cursor.execute("SELECT id_ESTADO, NOMBRE FROM ESTADO WHERE id_PAIS = ?", (id_pais,))
        rows = [{"id": r[0], "nombre": r[1]} for r in cursor.fetchall()]
        con.close()
        return jsonify(rows)
    except:
        return jsonify([])

@app.route("/api/ciudades/<int:id_estado>")
def api_ciudades(id_estado):
    try:
        con = conectar()
        cursor = con.cursor()
        cursor.execute("SELECT id_CIUDAD, NOMBRE FROM CIUDAD WHERE id_ESTADO = ?", (id_estado,))
        rows = [{"id": r[0], "nombre": r[1]} for r in cursor.fetchall()]
        con.close()
        return jsonify(rows)
    except:
        return jsonify([])

@app.route("/api/colonias/<int:id_ciudad>")
def api_colonias(id_ciudad):
    try:
        con = conectar()
        cursor = con.cursor()
        cursor.execute("SELECT id_COLONIA, NOMBRE, Codigo_Postal FROM COLONIA WHERE id_CIUDAD = ?", (id_ciudad,))
        rows = [{"id": r[0], "nombre": r[1], "cp": r[2]} for r in cursor.fetchall()]
        con.close()
        return jsonify(rows)
    except:
        return jsonify([])

# ═══════════════════════════════════════════
#  NUEVO USUARIO
# ═══════════════════════════════════════════
@app.route("/nuevo-usuario", methods=["GET", "POST"])
def nuevo_usuario():
    if "nombre" not in session:
        return redirect(url_for("login"))
    if session.get("rol") != "admin":
     flash("No tienes permiso para acceder aquí.", "error")
     return redirect(url_for("bienvenida"))

    if request.method == "POST":
        p1 = request.form.get("contrasena", "").strip()
        p2 = request.form.get("confirmar",  "").strip()
        if not p1:
            flash("La contraseña es obligatoria.", "error")
            return redirect(url_for("nuevo_usuario"))
        if p1 != p2:
            flash("Las contraseñas no coinciden.", "error")
            return redirect(url_for("nuevo_usuario"))
        try:
            con    = conectar()
            cursor = con.cursor()

            id_ciudad  = int(request.form.get("id_ciudad", 1))
            cp         = request.form.get("codigo_postal", "0").strip()
            cursor.execute("""
                INSERT INTO COLONIA (id_CIUDAD, Codigo_Postal, NOMBRE)
                VALUES (?, ?, ?)
            """, (id_ciudad, int(cp) if cp.isdigit() else 0,
                  request.form.get("colonia", "").strip() or "Sin colonia"))
            cursor.execute("SELECT @@IDENTITY")
            id_colonia = int(cursor.fetchone()[0])

            num_ext = request.form.get("num_ext", "").strip()
            cursor.execute("""
                INSERT INTO DIRECCION (id_COLONIA, CALLE, NUM_EXT, ENTRE_CALLE_1, ENTRE_CALLE_2)
                VALUES (?, ?, ?, ?, ?)
            """, (id_colonia,
                  request.form.get("calle", "").strip()         or None,
                  int(num_ext) if num_ext.isdigit()             else None,
                  request.form.get("entre_calle_1", "").strip() or None,
                  request.form.get("entre_calle_2", "").strip() or None))
            cursor.execute("SELECT @@IDENTITY")
            id_dir = int(cursor.fetchone()[0])

            cursor.execute("INSERT INTO LOGIN (CONTRASEÑA) VALUES (?)", (p1,))
            cursor.execute("SELECT @@IDENTITY")
            id_login = int(cursor.fetchone()[0])

            cursor.execute("""
                INSERT INTO USUARIOS
                (id_LOGIN, id_DIRECCION, id_SUCURSAL, NOMBRE, APELLID_P, APELLIDO_M, TELEFONO, CORREO,ROL)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?,?)
            """, (id_login, id_dir,
                  request.form.get("id_sucursal"),
                  request.form.get("nombre",     "").strip(),
                  request.form.get("apellido_p", "").strip(),
                  request.form.get("apellido_m", "").strip(),
                  request.form.get("telefono",   "").strip() or None,
                  request.form.get("correo",     "").strip(),
                  request.form.get("rol",        "vendedor"))) 
            con.commit()
            con.close()
            flash(f"Usuario '{request.form.get('nombre')}' creado correctamente. ✅", "success")
            return redirect(url_for("bienvenida"))
        except Exception as e:
            flash(f"Error al guardar: {e}", "error")

    try:
        con = conectar()
        cursor = con.cursor()
        cursor.execute("SELECT id_PAIS, NOMBRE FROM PAIS")
        paises = cursor.fetchall()
        cursor.execute("SELECT id_SUCURSAL, NOMBRE FROM SUCURSAL")
        sucursales = cursor.fetchall()
        con.close()
    except:
        paises = sucursales = []

    return render_template("nuevo_usuario.html", paises=paises, sucursales=sucursales)

# ═══════════════════════════════════════════
#  NUEVO ARTÍCULO
# ═══════════════════════════════════════════
@app.route("/nuevo-articulo", methods=["GET", "POST"])
def nuevo_articulo():
    if "nombre" not in session:
        return redirect(url_for("login"))
    if session.get("rol") != "admin":
        flash("No tienes permiso para acceder aquí.", "error")
        return redirect(url_for("bienvenida"))

    if request.method == "POST":
        try:
            con    = conectar()
            cursor = con.cursor()

            cursor.execute("""
                INSERT INTO PRODUCTO
                (id_SUCURSAL, id_TIPO, NOMBRE, MARCA, MODELO)
                VALUES (?, ?, ?, ?, ?)
            """, (
                request.form.get("id_sucursal"),
                request.form.get("id_tipo"),
                request.form.get("nombre",  "").strip(),
                request.form.get("marca",   "").strip(),
                request.form.get("modelo",  "").strip() or None
            ))
            cursor.execute("SELECT @@IDENTITY")
            id_producto = int(cursor.fetchone()[0])

            indices = request.form.getlist("variante_idx")
            codigos_generados = []

            for i in indices:
                codigo = request.form.get(f"codigo_{i}", "").strip() or generar_sku()
                codigos_generados.append(codigo)
                cursor.execute("""
                    INSERT INTO VARIANTE
                    (id_PRODUCTO, TALLA, COLOR, CODIGO_BARRAS, STOCK, PRECIO)
                    VALUES (?, ?, ?, ?, ?, ?)
                """, (
                    id_producto,
                    request.form.get(f"talla_{i}", "").strip() or None,
                    request.form.get(f"color_{i}", "").strip() or None,
                    codigo,
                    int(request.form.get(f"stock_{i}", 0)),
                    request.form.get(f"precio_{i}", "").strip() or None
                ))

            con.commit()
            con.close()
            flash(f"Producto guardado con {len(indices)} variante(s). ✅", "success")
            return redirect(url_for("bienvenida"))

        except Exception as e:
            manejar_error(e, "No se pudo guardar el producto. Intenta de nuevo.")

    try:
        con = conectar()
        cursor = con.cursor()
        cursor.execute("SELECT id_SUCURSAL, NOMBRE FROM SUCURSAL")
        sucursales = cursor.fetchall()
        cursor.execute("SELECT id_TIPO, DESCRIPCION, CAST(TIENE_VARIANTES AS INT) FROM TIPO")
        tipos = cursor.fetchall()
        con.close()
    except:
        sucursales = tipos = []

    return render_template("nuevo_articulo.html", sucursales=sucursales, tipos=tipos)

# ═══════════════════════════════════════════
#  ETIQUETAS
# ═══════════════════════════════════════════
@app.route("/etiquetas")
def etiquetas():
    if "nombre" not in session:
        return redirect(url_for("login"))
    if session.get("rol") != "admin":
        flash("No tienes permiso para acceder aquí.", "error")
        return redirect(url_for("bienvenida"))
    try:
        con = conectar()
        cursor = con.cursor()
        cursor.execute("""
            SELECT V.id_VARIANTE, P.NOMBRE, P.MARCA, V.TALLA, V.COLOR,
            V.PRECIO, V.CODIGO_BARRAS, V.STOCK, V.ETIQUETA_IMPRESA
            FROM VARIANTE V
            INNER JOIN PRODUCTO P ON V.id_PRODUCTO = P.id_PRODUCTO
            WHERE V.STOCK > 0
            ORDER BY V.ETIQUETA_IMPRESA ASC, P.NOMBRE
            """)
        variantes = cursor.fetchall()
        con.close()
    except:
        variantes = []
    return render_template("etiquetas.html", variantes=variantes)

@app.route("/generar-etiquetas", methods=["POST"])
def generar_etiquetas():
    if "nombre" not in session:
        return redirect(url_for("login"))
    if session.get("rol") != "admin":
        flash("No tienes permiso para acceder aquí.", "error")
        return redirect(url_for("bienvenida"))

    ids = request.form.getlist("variantes")
    if not ids:
        flash("Selecciona al menos una variante.", "error")
        return redirect(url_for("etiquetas"))

    try:
        con = conectar()
        cursor = con.cursor()
        placeholders = ",".join(["?" for _ in ids])
        cursor.execute(f"""
            SELECT P.NOMBRE, P.MARCA, V.TALLA, V.COLOR, V.PRECIO, V.CODIGO_BARRAS
            FROM VARIANTE V
            INNER JOIN PRODUCTO P ON V.id_PRODUCTO = P.id_PRODUCTO
            WHERE V.id_VARIANTE IN ({placeholders})
        """, ids)
        variantes = cursor.fetchall()
        pdf_path = generar_pdf_etiquetas(variantes)
        # Marcar como impresas
        placeholders = ",".join(["?" for _ in ids])
        cursor.execute(f"""
                        UPDATE VARIANTE SET ETIQUETA_IMPRESA = 1
                        WHERE id_VARIANTE IN ({placeholders})
                        """, ids)
        con.commit()
        con.close()
        return send_file(pdf_path, as_attachment=False, mimetype="application/pdf")
    except Exception as e:
        manejar_error(e, "No se pudieron generar las etiquetas.")
        return redirect(url_for("etiquetas"))

def generar_pdf_etiquetas(variantes):
    import barcode
    from barcode.writer import ImageWriter
    from reportlab.pdfgen import canvas
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.units import cm
    import tempfile, os

    ancho_hoja, alto_hoja = letter
    cols      = 3
    filas     = 4
    ancho_etq = ancho_hoja / cols
    alto_etq  = alto_hoja  / filas
    margen    = 0.3 * cm

    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".pdf")
    pdf = canvas.Canvas(tmp.name, pagesize=letter)

    for i, var in enumerate(variantes):
        nombre, marca, talla, color, precio, codigo = var
        col  = i % cols
        fila = (i // cols) % filas
        if i > 0 and i % (cols * filas) == 0:
            pdf.showPage()

        x = col  * ancho_etq + margen
        y = alto_hoja - (fila + 1) * alto_etq + margen

        pdf.setStrokeColorRGB(0.6, 0.1, 0.1)
        pdf.setLineWidth(0.5)
        pdf.rect(col * ancho_etq, alto_hoja - (fila + 1) * alto_etq, ancho_etq, alto_etq)

        pdf.setFillColorRGB(0.1, 0.04, 0)
        pdf.setFont("Helvetica-Bold", 9)
        pdf.drawCentredString(col * ancho_etq + ancho_etq / 2, y + alto_etq - 1.2 * cm, "PAHJAL Western Boots")

        pdf.setFont("Helvetica-Bold", 8)
        pdf.setFillColorRGB(0, 0, 0)
        pdf.drawCentredString(col * ancho_etq + ancho_etq / 2, y + alto_etq - 1.9 * cm, (nombre or "")[:22])

        pdf.setFont("Helvetica", 7)
        pdf.setFillColorRGB(0.3, 0.3, 0.3)
        pdf.drawCentredString(col * ancho_etq + ancho_etq / 2, y + alto_etq - 2.4 * cm, marca or "")

        # Talla y color
        detalle = []
        if talla: detalle.append(f"Talla: {talla}")
        if color: detalle.append(f"Color: {color}")
        if detalle:
            pdf.setFont("Helvetica", 7)
            pdf.setFillColorRGB(0.2, 0.2, 0.2)
            pdf.drawCentredString(col * ancho_etq + ancho_etq / 2, y + alto_etq - 2.85 * cm,
                                  " | ".join(detalle))

        try:
            codigo_str = str(codigo) if codigo else "PAH-0000"
            BC  = barcode.get_barcode_class("code128")
            bc  = BC(codigo_str, writer=ImageWriter())
            tmp_bc = tempfile.NamedTemporaryFile(delete=False, suffix=".png")
            bc.write(tmp_bc, options={"module_width": 0.4, "module_height": 8,
                                      "quiet_zone": 2, "font_size": 6,
                                      "text_distance": 2, "write_text": False})
            tmp_bc.close()
            pdf.drawImage(tmp_bc.name,
                          col * ancho_etq + margen, y + alto_etq - 4.7 * cm,
                          width=ancho_etq - (margen * 2), height=1.5 * cm)
            os.unlink(tmp_bc.name)
        except:
            pass

        pdf.setFont("Helvetica", 6)
        pdf.setFillColorRGB(0, 0, 0)
        pdf.drawCentredString(col * ancho_etq + ancho_etq / 2, y + alto_etq - 5.2 * cm, str(codigo or ""))

        pdf.setFont("Helvetica-Bold", 10)
        pdf.setFillColorRGB(0.6, 0.1, 0.1)
        precio_txt = f"${precio:,.2f}" if precio else "Sin precio"
        pdf.drawCentredString(col * ancho_etq + ancho_etq / 2, y + alto_etq - 5.9 * cm, precio_txt)

    pdf.save()
    return tmp.name

# ═══════════════════════════════════════════
#  NUEVA VENTA
# ═══════════════════════════════════════════
@app.route("/nueva-venta", methods=["GET", "POST"])
def nueva_venta():
    if "nombre" not in session:
        return redirect(url_for("login"))

    if "carrito" not in session:
        session["carrito"] = []

    try:
        con = conectar()
        cursor = con.cursor()
        cursor.execute("SELECT id_CLIENTE, NOMBRE, APELLIDO_P FROM CLIENTE")
        clientes = cursor.fetchall()
        cursor.execute("SELECT id_METODO_PAGO, DESCRIPCION FROM METODO_PAGO")
        metodos = cursor.fetchall()
        con.close()
    except:
        clientes = metodos = []

    return render_template("nueva_venta.html",
                           carrito=session["carrito"],
                           metodos=metodos)

@app.route("/api/buscar-articulo")
def buscar_articulo():
    termino = request.args.get("q", "").strip()
    if not termino:
        return jsonify([])
    try:
        con = conectar()
        cursor = con.cursor()
        cursor.execute("""
            SELECT V.id_VARIANTE, P.NOMBRE, P.MARCA, V.TALLA, V.COLOR,
                   V.PRECIO, V.CODIGO_BARRAS, V.STOCK
            FROM VARIANTE V
            INNER JOIN PRODUCTO P ON V.id_PRODUCTO = P.id_PRODUCTO
            WHERE (V.CODIGO_BARRAS = ? OR P.NOMBRE LIKE ?)
            AND V.STOCK > 0
        """, (termino, f"%{termino}%"))
        rows = [{
            "id":     str(r[0]),
            "nombre": r[1],
            "marca":  r[2],
            "talla":  r[3] or "—",
            "color":  r[4] or "—",
            "precio": float(r[5]) if r[5] else 0,
            "codigo": r[6],
            "stock":  r[7]
        } for r in cursor.fetchall()]
        con.close()
        return jsonify(rows)
    except:
        return jsonify([])

@app.route("/carrito/agregar", methods=["POST"])
def agregar_carrito():
    id_var  = request.form.get("id_articulo")
    nombre  = request.form.get("nombre")
    precio  = float(request.form.get("precio", 0))
    codigo  = request.form.get("codigo")
    talla   = request.form.get("talla", "—")
    color   = request.form.get("color", "—")
    carrito = session.get("carrito", [])

    # Consultar stock real de la BD
    try:
        con = conectar()
        cursor = con.cursor()
        cursor.execute("SELECT STOCK FROM VARIANTE WHERE id_VARIANTE = ?", (id_var,))
        row = cursor.fetchone()
        con.close()
        stock_disponible = row[0] if row else 0
    except:
        stock_disponible = 0

    # Calcular cuánto ya hay en el carrito
    cantidad_en_carrito = 0
    for item in carrito:
        if item["id"] == id_var:
            cantidad_en_carrito = item["cantidad"]
            break

    # Validar si hay stock suficiente
    if cantidad_en_carrito >= stock_disponible:
        flash(f"No hay más stock disponible de {nombre}. Solo hay {stock_disponible} pieza(s).", "error")
        return redirect(url_for("nueva_venta"))

    # Agregar o incrementar en carrito
    for item in carrito:
        if item["id"] == id_var:
            item["cantidad"] += 1
            item["subtotal"]  = round(item["cantidad"] * item["precio"], 2)
            session["carrito"] = carrito
            session.modified   = True
            return redirect(url_for("nueva_venta"))

    carrito.append({
        "id":       id_var,
        "nombre":   nombre,
        "precio":   precio,
        "codigo":   codigo,
        "talla":    talla,
        "color":    color,
        "cantidad": 1,
        "subtotal": precio
    })
    session["carrito"] = carrito
    session.modified   = True
    return redirect(url_for("nueva_venta"))

@app.route("/carrito/quitar/<id_art>")
def quitar_carrito(id_art):
    session["carrito"] = [i for i in session.get("carrito", []) if i["id"] != id_art]
    session.modified   = True
    return redirect(url_for("nueva_venta"))

@app.route("/carrito/limpiar")
def limpiar_carrito():
    session["carrito"] = []
    session.modified   = True
    return redirect(url_for("nueva_venta"))

@app.route("/cobrar", methods=["POST"])
def cobrar():
    if "nombre" not in session:
        return redirect(url_for("login"))

    carrito = session.get("carrito", [])
    if not carrito:
        flash("El carrito está vacío.", "error")
        return redirect(url_for("nueva_venta"))

    id_cliente  = 2
    id_metodo   = request.form.get("id_metodo")
    id_usuario  = session.get("id_usuario")
    id_sucursal = session.get("id_sucursal", 1)

    if not id_cliente or not id_metodo:
        flash("Selecciona el método de pago.", "error")
        return redirect(url_for("nueva_venta"))

    try:
        con    = conectar()
        cursor = con.cursor()

        cursor.execute("""
            INSERT INTO TICKET
            (id_METODO_PAGO, id_CLIENTE, id_USUARIOS, id_SUCURSAL, FECHA, HORA)
            VALUES (?, ?, ?, ?, CAST(GETDATE() AS DATE), CAST(GETDATE() AS TIME))
        """, (id_metodo, id_cliente, id_usuario, id_sucursal))
        cursor.execute("SELECT @@IDENTITY")
        id_ticket = int(cursor.fetchone()[0])

        for item in carrito:
            cursor.execute("""
                INSERT INTO CANTIDAD_ARTICULOS (id_ARTICULO, id_TICKET, CANTIDAD)
                VALUES (?, ?, ?)
            """, (item["id"], id_ticket, item["cantidad"]))
            cursor.execute("""
        UPDATE VARIANTE
        SET STOCK = STOCK - ?
        WHERE id_VARIANTE = ?
    """, (item["cantidad"], item["id"]))

        con.commit()

        cursor.execute("""
            SELECT T.id_TICKET, T.FECHA, T.HORA,
                   C.NOMBRE + ' ' + C.APELLIDO_P  AS CLIENTE,
                   U.NOMBRE + ' ' + U.APELLID_P   AS VENDEDOR,
                   M.DESCRIPCION                   AS METODO_PAGO,
                   S.NOMBRE                        AS SUCURSAL
            FROM TICKET T
            INNER JOIN CLIENTE     C ON T.id_CLIENTE     = C.id_CLIENTE
            INNER JOIN USUARIOS    U ON T.id_USUARIOS    = U.id_USUARIOS
            INNER JOIN METODO_PAGO M ON T.id_METODO_PAGO = M.id_METODO_PAGO
            INNER JOIN SUCURSAL    S ON T.id_SUCURSAL    = S.id_SUCURSAL
            WHERE T.id_TICKET = ?
        """, (id_ticket,))
        ticket = cursor.fetchone()
        con.close()

        total           = sum(i["subtotal"] for i in carrito)
        carrito_copia   = list(carrito)
        session["carrito"] = []
        session.modified   = True

        return render_template("ticket_imprimible.html",
                               ticket=ticket, carrito=carrito_copia, total=total)
    except Exception as e:
        manejar_error(e, "No se pudo procesar la venta. Intenta de nuevo.")
        return redirect(url_for("nueva_venta"))
    
# ═══════════════════════════════════════════
#  CORTE DEL DÍA
# ═══════════════════════════════════════════
@app.route("/corte")
def corte():

    if "nombre" not in session:
        return redirect(url_for("login"))
    if session.get("rol") != "admin":
        flash("No tienes permiso para acceder aquí.", "error")
        return redirect(url_for("bienvenida"))

    fecha = request.args.get("fecha")

    try:
        con = conectar()
        cursor = con.cursor()

        if fecha:
            cursor.execute("""
                SELECT T.id_TICKET, T.FECHA, T.HORA,
                       U.NOMBRE + ' ' + U.APELLID_P AS VENDEDOR,
                       M.DESCRIPCION AS METODO_PAGO,
                       SUM(V.PRECIO * CA.CANTIDAD) AS TOTAL
                FROM TICKET T
                INNER JOIN USUARIOS U          ON T.id_USUARIOS    = U.id_USUARIOS
                INNER JOIN METODO_PAGO M       ON T.id_METODO_PAGO = M.id_METODO_PAGO
                INNER JOIN CANTIDAD_ARTICULOS CA ON T.id_TICKET     = CA.id_TICKET
                INNER JOIN VARIANTE V          ON CA.id_ARTICULO    = V.id_VARIANTE
                WHERE T.FECHA = ?
                GROUP BY T.id_TICKET, T.FECHA, T.HORA, U.NOMBRE, U.APELLID_P, M.DESCRIPCION
                ORDER BY T.HORA
            """, (fecha,))
        else:
            cursor.execute("""
                SELECT T.id_TICKET, T.FECHA, T.HORA,
                       U.NOMBRE + ' ' + U.APELLID_P AS VENDEDOR,
                       M.DESCRIPCION AS METODO_PAGO,
                       SUM(V.PRECIO * CA.CANTIDAD) AS TOTAL
                FROM TICKET T
                INNER JOIN USUARIOS U          ON T.id_USUARIOS    = U.id_USUARIOS
                INNER JOIN METODO_PAGO M       ON T.id_METODO_PAGO = M.id_METODO_PAGO
                INNER JOIN CANTIDAD_ARTICULOS CA ON T.id_TICKET     = CA.id_TICKET
                INNER JOIN VARIANTE V          ON CA.id_ARTICULO    = V.id_VARIANTE
                WHERE T.FECHA = CAST(GETDATE() AS DATE)
                GROUP BY T.id_TICKET, T.FECHA, T.HORA, U.NOMBRE, U.APELLID_P, M.DESCRIPCION
                ORDER BY T.HORA
            """)

        tickets = cursor.fetchall()

        # Desglose por método de pago
        cursor.execute("""
            SELECT M.DESCRIPCION, SUM(V.PRECIO * CA.CANTIDAD) AS TOTAL
            FROM TICKET T
            INNER JOIN METODO_PAGO M       ON T.id_METODO_PAGO = M.id_METODO_PAGO
            INNER JOIN CANTIDAD_ARTICULOS CA ON T.id_TICKET     = CA.id_TICKET
            INNER JOIN VARIANTE V          ON CA.id_ARTICULO    = V.id_VARIANTE
            WHERE T.FECHA = {}
            GROUP BY M.DESCRIPCION
        """.format("?" if fecha else "CAST(GETDATE() AS DATE)"),
        (fecha,) if fecha else ())

        desglose = cursor.fetchall()
        con.close()

        total_dia = sum(t[5] for t in tickets) if tickets else 0

    except Exception as e:
        manejar_error(e, "No se pudo generar el corte del día.")
        tickets  = []
        desglose = []
        total_dia = 0

    return render_template("corte.html", tickets=tickets,
                           total_dia=total_dia, fecha=fecha,
                           desglose=desglose)

@app.route("/articulos")
def articulos():
    if "nombre" not in session:
        return redirect(url_for("login"))
    if session.get("rol") != "admin":
        flash("No tienes permiso para acceder aquí.", "error")
        return redirect(url_for("bienvenida"))
    try:
        con = conectar()
        cursor = con.cursor()
        cursor.execute("""
            SELECT P.id_PRODUCTO, P.NOMBRE, P.MARCA, P.MODELO,
                   T.DESCRIPCION,
                   COUNT(V.id_VARIANTE) AS NUM_VARIANTES,
                   SUM(V.STOCK) AS STOCK_TOTAL
            FROM PRODUCTO P
            INNER JOIN TIPO T ON P.id_TIPO = T.id_TIPO
            LEFT JOIN VARIANTE V ON P.id_PRODUCTO = V.id_PRODUCTO
            GROUP BY P.id_PRODUCTO, P.NOMBRE, P.MARCA, P.MODELO, T.DESCRIPCION
            ORDER BY P.NOMBRE
        """)
        productos = cursor.fetchall()
        con.close()
    except Exception as e:
        flash(f"Error: {e}", "error")
        productos = []
    return render_template("articulos.html", productos=productos)


@app.route("/articulos/<int:id_producto>")
def detalle_producto(id_producto):
    if "nombre" not in session:
        return redirect(url_for("login"))
    if session.get("rol") != "admin":
        flash("No tienes permiso para acceder aquí.", "error")
        return redirect(url_for("bienvenida"))
    try:
        con = conectar()
        cursor = con.cursor()
        cursor.execute("""
            SELECT id_PRODUCTO, NOMBRE, MARCA, MODELO FROM PRODUCTO
            WHERE id_PRODUCTO = ?
        """, (id_producto,))
        producto = cursor.fetchone()

        cursor.execute("""
            SELECT id_VARIANTE, TALLA, COLOR, CODIGO_BARRAS, STOCK, PRECIO
            FROM VARIANTE WHERE id_PRODUCTO = ?
            ORDER BY TALLA, COLOR
        """, (id_producto,))
        variantes = cursor.fetchall()
        con.close()
    except Exception as e:
        flash(f"Error: {e}", "error")
        producto = None
        variantes = []
    return render_template("detalle_producto.html",
                           producto=producto, variantes=variantes)


@app.route("/articulos/<int:id_producto>/agregar-variante", methods=["POST"])
def agregar_variante(id_producto):
    if "nombre" not in session:
        return redirect(url_for("login"))
    if session.get("rol") != "admin":
        flash("No tienes permiso.", "error")
        return redirect(url_for("bienvenida"))
    try:
        con = conectar()
        cursor = con.cursor()
        codigo = request.form.get("codigo", "").strip() or generar_sku()
        cursor.execute("""
            INSERT INTO VARIANTE (id_PRODUCTO, TALLA, COLOR, CODIGO_BARRAS, STOCK, PRECIO)
            VALUES (?, ?, ?, ?, ?, ?)
        """, (
            id_producto,
            request.form.get("talla",  "").strip() or None,
            request.form.get("color",  "").strip() or None,
            codigo,
            int(request.form.get("stock", 0)),
            request.form.get("precio", "").strip() or None
        ))
        con.commit()
        con.close()
        flash("Variante agregada. ✅", "success")
    except Exception as e:
        flash(f"Error: {e}", "error")
    return redirect(url_for("detalle_producto", id_producto=id_producto))


@app.route("/variante/<int:id_variante>/editar", methods=["POST"])
def editar_variante(id_variante):

    if "nombre" not in session:
        return redirect(url_for("login"))
    if session.get("rol") != "admin":
        flash("No tienes permiso.", "error")
        return redirect(url_for("bienvenida"))
    try:
        con = conectar()
        cursor = con.cursor()
        cursor.execute("""
            UPDATE VARIANTE SET TALLA=?, COLOR=?, CODIGO_BARRAS=?, STOCK=?, PRECIO=?
            WHERE id_VARIANTE=?
        """, (
            request.form.get("talla",  "").strip() or None,
            request.form.get("color",  "").strip() or None,
            request.form.get("codigo", "").strip() or None,
            int(request.form.get("stock", 0)),
            request.form.get("precio", "").strip() or None,
            id_variante
        ))
        id_producto = request.form.get("id_producto")
        con.commit()
        con.close()
        flash("Variante actualizada. ✅", "success")
    except Exception as e:
        manejar_error(e, "No se pudo completar la acción.")
        id_producto = request.form.get("id_producto")
    return redirect(url_for("detalle_producto", id_producto=id_producto))

@app.route("/variante/<int:id_variante>/eliminar", methods=["POST"])
def eliminar_variante(id_variante):
    if "nombre" not in session:
        return redirect(url_for("login"))
    if session.get("rol") != "admin":
        flash("No tienes permiso.", "error")
        return redirect(url_for("bienvenida"))

    id_producto = request.form.get("id_producto")
    forzar      = request.form.get("forzar") == "1"

    try:
        con    = conectar()
        cursor = con.cursor()

        # Verificar si tiene ventas registradas
        cursor.execute("""
            SELECT COUNT(*) FROM CANTIDAD_ARTICULOS WHERE id_ARTICULO = ?
        """, (id_variante,))
        tiene_ventas = cursor.fetchone()[0] > 0

        if tiene_ventas and not forzar:
            con.close()
            flash("Esta variante tiene ventas registradas. Confirma para eliminarla de todas formas.", "error")
            return redirect(url_for("detalle_producto", id_producto=id_producto,
                                    confirmar_eliminar=id_variante))

        if tiene_ventas and forzar:
            # Eliminar primero las referencias en CANTIDAD_ARTICULOS
            cursor.execute("DELETE FROM CANTIDAD_ARTICULOS WHERE id_ARTICULO = ?", (id_variante,))

        cursor.execute("DELETE FROM VARIANTE WHERE id_VARIANTE = ?", (id_variante,))
        con.commit()
        con.close()
        flash("Variante eliminada. ✅", "success")
    except Exception as e:
        manejar_error(e, "No se pudo completar la acción.")

    return redirect(url_for("detalle_producto", id_producto=id_producto))

@app.route("/buscar-inventario")
def buscar_inventario():
    if "nombre" not in session:
        return redirect(url_for("login"))

    nombre = request.args.get("nombre", "").strip()
    marca  = request.args.get("marca",  "").strip()
    modelo = request.args.get("modelo", "").strip()
    talla  = request.args.get("talla",  "").strip()
    color  = request.args.get("color",  "").strip()

    resultados = []
    buscando   = any([nombre, marca, modelo, talla, color])

    if buscando:
        try:
            con    = conectar()
            cursor = con.cursor()

            query = """
                SELECT P.NOMBRE, P.MARCA, P.MODELO,
                       V.TALLA, V.COLOR, V.CODIGO_BARRAS,
                       CASE WHEN V.STOCK > 0 THEN 1 ELSE 0 END AS DISPONIBLE
                FROM VARIANTE V
                INNER JOIN PRODUCTO P ON V.id_PRODUCTO = P.id_PRODUCTO
                WHERE 1=1
            """
            params = []

            if nombre:
                query += " AND P.NOMBRE LIKE ?"
                params.append(f"%{nombre}%")
            if marca:
                query += " AND P.MARCA LIKE ?"
                params.append(f"%{marca}%")
            if modelo:
                query += " AND P.MODELO LIKE ?"
                params.append(f"%{modelo}%")
            if talla:
                query += " AND V.TALLA LIKE ?"
                params.append(f"%{talla}%")
            if color:
                query += " AND V.COLOR LIKE ?"
                params.append(f"%{color}%")

            query += " ORDER BY P.NOMBRE, V.TALLA, V.COLOR"

            cursor.execute(query, params)
            resultados = cursor.fetchall()
            con.close()
        except Exception as e:
            flash(f"Error: {e}", "error")

    return render_template("buscar_inventario.html",
                           resultados=resultados,
                           buscando=buscando,
                           nombre=nombre, marca=marca,
                           modelo=modelo, talla=talla, color=color)

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from flask import send_file
import tempfile, os

@app.route("/descargar-excel/<int:id_ticket>")
def descargar_excel(id_ticket):

    if "nombre" not in session:
        return redirect(url_for("login"))
    try:
        con    = conectar()
        cursor = con.cursor()

        # Datos del ticket
        cursor.execute("""
            SELECT T.id_TICKET, T.FECHA, T.HORA,
                   C.NOMBRE + ' ' + C.APELLIDO_P  AS CLIENTE,
                   U.NOMBRE + ' ' + U.APELLID_P   AS VENDEDOR,
                   M.DESCRIPCION                   AS METODO_PAGO,
                   S.NOMBRE                        AS SUCURSAL
            FROM TICKET T
            INNER JOIN CLIENTE     C ON T.id_CLIENTE     = C.id_CLIENTE
            INNER JOIN USUARIOS    U ON T.id_USUARIOS    = U.id_USUARIOS
            INNER JOIN METODO_PAGO M ON T.id_METODO_PAGO = M.id_METODO_PAGO
            INNER JOIN SUCURSAL    S ON T.id_SUCURSAL    = S.id_SUCURSAL
            WHERE T.id_TICKET = ?
        """, (id_ticket,))
        ticket = cursor.fetchone()

        # Artículos del ticket
        cursor.execute("""
            SELECT P.NOMBRE, P.MARCA, V.TALLA, V.COLOR,
                   V.CODIGO_BARRAS, CA.CANTIDAD, V.PRECIO,
                   (CA.CANTIDAD * V.PRECIO) AS SUBTOTAL
            FROM CANTIDAD_ARTICULOS CA
            INNER JOIN VARIANTE V ON CA.id_ARTICULO = V.id_VARIANTE
            INNER JOIN PRODUCTO P ON V.id_PRODUCTO  = P.id_PRODUCTO
            WHERE CA.id_TICKET = ?
        """, (id_ticket,))
        items = cursor.fetchall()
        con.close()

        # Crear Excel
        wb = Workbook()
        ws = wb.active
        ws.title = f"Ticket {id_ticket}"

        # Estilos
        rojo    = PatternFill("solid", fgColor="CC0000")
        verde   = PatternFill("solid", fgColor="217346")
        blanco  = Font(color="FFFFFF", bold=True)
        centrar = Alignment(horizontal="center")

        # Encabezado tienda
        ws.merge_cells("A1:H1")
        ws["A1"] = "PAHJAL WESTERN BOOTS"
        ws["A1"].font      = Font(bold=True, size=14, color="CC0000")
        ws["A1"].alignment = centrar

        ws.merge_cells("A2:H2")
        ws["A2"] = f"Ticket #{ticket[0]}  |  {ticket[1]}  {ticket[2]}"
        ws["A2"].alignment = centrar
        ws["A2"].font      = Font(size=10)

        ws.merge_cells("A3:H3")
        ws["A3"] = f"Sucursal: {ticket[6]}  |  Vendedor: {ticket[4]}  |  Pago: {ticket[5]}"
        ws["A3"].alignment = centrar
        ws["A3"].font      = Font(size=9, italic=True)

        ws.append([])  # fila vacía

        # Encabezados tabla
        headers = ["NOMBRE", "MARCA", "TALLA", "COLOR", "CÓDIGO", "CANTIDAD", "PRECIO", "SUBTOTAL"]
        ws.append(headers)
        for col in range(1, 9):
            cell = ws.cell(row=5, column=col)
            cell.fill      = rojo
            cell.font      = blanco
            cell.alignment = centrar

        # Datos
        total = 0
        for item in items:
            ws.append([
                item[0],           # nombre
                item[1],           # marca
                item[2] or "—",    # talla
                item[3] or "—",    # color
                item[4] or "—",    # codigo
                item[5],           # cantidad
                float(item[6]) if item[6] else 0,   # precio
                float(item[7]) if item[7] else 0,   # subtotal
            ])
            total += float(item[7]) if item[7] else 0

        # Fila total
        fila_total = ws.max_row + 1
        ws.cell(row=fila_total, column=7, value="TOTAL:")
        ws.cell(row=fila_total, column=7).font      = Font(bold=True)
        ws.cell(row=fila_total, column=7).alignment = Alignment(horizontal="right")
        ws.cell(row=fila_total, column=8, value=round(total, 2))
        ws.cell(row=fila_total, column=8).font = Font(bold=True, color="CC0000")

        # Ancho de columnas
        # Ancho de columnas
        anchos = {"A":22, "B":15, "C":8, "D":12, "E":15, "F":10, "G":10, "H":12}
        for col, ancho in anchos.items():
            ws.column_dimensions[col].width = ancho

        # Guardar temporal
        tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")
        wb.save(tmp.name)
        tmp.close()

        return send_file(tmp.name, as_attachment=True,
                         download_name=f"Ticket_{id_ticket}_PAHJAL.xlsx",
                         mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    except Exception as e:
        flash(f"Error al generar Excel: {e}", "error")
        return redirect(url_for("nueva_venta"))
    
@app.route("/service-worker.js")
def service_worker():
    return send_file("service-worker.js", mimetype="application/javascript")

# ═══════════════════════════════════════════
#  LOGOUT
# ═══════════════════════════════════════════
@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))


if __name__ == "__main__":
    app.run(debug=True)
